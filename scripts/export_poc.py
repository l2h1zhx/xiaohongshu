"""将私有PoC结果导出为核对表，不冒充月结或24h历史数据。"""

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from koc_insight.report import display_time, text
from koc_insight.urls import validate_url


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--links", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.exists():
        raise ValueError("不覆盖已有核对表")
    result = json.loads(args.result.read_text())
    links = json.loads(args.links.read_text())
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "真实笔记核对"
    headers = [
        "样本编号",
        "品牌",
        "笔记ID",
        "作者ID",
        "作者昵称",
        "标题",
        "正文",
        "话题",
        "类型",
        "发布时间(北京时间)",
        "本次采集时间(北京时间)",
        "点赞",
        "收藏",
        "评论",
        "分享",
        "曝光",
        "阅读",
        "原始指标",
        "精度与缺失说明",
        "数据来源",
        "结果",
        "查看原笔记",
    ]
    sheet.append(headers)
    for sample in result["samples"]:
        detail = sample.get("detail", {})
        counts = detail.get("counts", {})
        url = validate_url(links[sample["sample"] - 1])
        values = [
            sample["sample"],
            "待确认（授权PoC）",
            detail.get("note_id"),
            detail.get("author_id"),
            detail.get("author_name"),
            detail.get("title"),
            detail.get("description"),
            detail.get("tags"),
            detail.get("note_type"),
            display_time(detail.get("published_at")),
            display_time(sample.get("captured_at")),
        ]
        values += [
            counts.get(k)
            for k in ("likes", "collects", "comments", "shares", "exposure", "reads")
        ]
        values += [
            json.dumps(detail.get("raw_counts", {}), ensure_ascii=False),
            json.dumps(detail.get("availability", {}), ensure_ascii=False),
            detail.get("source"),
            sample["status"],
            url,
        ]
        sheet.append([text(v) for v in values])
        linkcell = sheet.cell(sheet.max_row, sheet.max_column)
        linkcell.hyperlink = url
        linkcell.style = "Hyperlink"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174E63")
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
    info = workbook.create_sheet("采集说明")
    info.append(["项目", "说明"])
    info.append(
        ["用途", "用户授权链接的真实详情PoC核对；不是月结数据，也不是过去24h数据。"]
    )
    info.append(
        ["来源", f"XHS-Downloader {result['upstream_sha']} 的真实 /xhs/detail 路由"]
    )
    info.append(
        [
            "安全",
            f"系统CA证书校验；download=false；新增媒体文件={result['new_media_files']}",
        ]
    )
    info.append(
        [
            "缺失指标",
            "未知值保持空白，原始文本和解析精度列供核对；曝光和阅读未由公开详情取得。",
        ]
    )
    info.append(["检查范围", "已读取标题、正文、话题等；未核验图片/视频画面。"])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 110
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    args.output.chmod(0o600)
    print(
        json.dumps(
            {"path": str(args.output), "real_notes": result["successful_notes"]},
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
