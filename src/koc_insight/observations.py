"""品牌自然笔记：仅作观察，不进入投放、月报或后续自动采集。"""

import hashlib
import json
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from .domain import SH
from .feishu import Feishu
from .project_import import import_plan

TABLE_NAME = "09 品牌自然观察"
SOURCE_LABEL = "vishin_wanxie_filtered_notes.xlsx"
FIELDS = {
    "笔记标题": 1,
    "Note ID": 1,
    "观察状态": 3,
    "关联说明": 1,
    "命中关键词": 1,
    "作者昵称": 1,
    "发布时间": 5,
    "笔记类型": 1,
    "正文摘要": 1,
    "话题": 1,
    "点赞": 2,
    "收藏": 2,
    "评论": 2,
    "分享": 2,
    "数据口径": 1,
    "原笔记链接": 15,
    "来源文件": 1,
    "来源键": 1,
}


def _checked_workbook(path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("需要不超过20MB的 .xlsx 观察文件")
    with ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("观察文件展开超过100MB")
    return path


def _timestamp(value):
    instant = (
        value
        if isinstance(value, datetime)
        else datetime.strptime(str(value), "%Y-%m-%d %H:%M:%S")
    )
    if instant.tzinfo is None:
        instant = instant.replace(tzinfo=SH)
    return int(instant.timestamp() * 1000)


def _number(value):
    return None if value in (None, "") else int(value)


def _records(sheet, status):
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = {name: index for index, name in enumerate(rows[0]) if name}
    required = {"筛选结论", "笔记ID", "标题", "发布时间（北京时间）", "笔记链接URL"}
    if not required.issubset(headers):
        raise ValueError("观察文件缺少必要列")
    records = []

    def value(row, name):
        index = headers.get(name)
        return row[index] if index is not None else None

    for row in rows[1:]:
        note_id = row[headers["笔记ID"]]
        if not note_id:
            continue
        note_id = str(note_id)
        if len(note_id) != 24 or any(
            char not in "0123456789abcdef" for char in note_id
        ):
            raise ValueError("观察文件包含无效 Note ID")
        link = str(row[headers["笔记链接URL"]] or "")
        if not link.startswith("https://www.xiaohongshu.com/explore/"):
            raise ValueError("观察文件包含无效笔记链接")
        record = {
            "笔记标题": str(row[headers["标题"]] or "无标题"),
            "Note ID": note_id,
            "观察状态": status,
            "关联说明": str(row[headers["筛选结论"]] or ""),
            "命中关键词": str(value(row, "命中关键词") or ""),
            "作者昵称": str(value(row, "作者昵称") or ""),
            "发布时间": _timestamp(row[headers["发布时间（北京时间）"]]),
            "笔记类型": str(value(row, "笔记类型") or ""),
            "正文摘要": str(value(row, "正文/描述") or ""),
            "话题": str(value(row, "标签列表") or ""),
            "数据口径": "来源文件导入时的单次公开数据；不自动刷新，不纳入项目交付",
            "原笔记链接": {"text": "查看原笔记", "link": link},
            "来源文件": SOURCE_LABEL,
            "来源键": "natural:" + note_id,
        }
        for source, target in (
            ("点赞数", "点赞"),
            ("收藏数", "收藏"),
            ("评论数", "评论"),
            ("分享数", "分享"),
        ):
            number = _number(value(row, source))
            if number is not None:
                record[target] = number
        records.append(record)
    return records


def observation_plan(path):
    path = _checked_workbook(path)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if not {"相关笔记", "待复核"}.issubset(book.sheetnames):
            raise ValueError("观察文件需要包含“相关笔记”和“待复核”工作表")
        records = _records(book["相关笔记"], "相关") + _records(
            book["待复核"], "待复核"
        )
    finally:
        book.close()
    keys = [record["来源键"] for record in records]
    if len(keys) != len(set(keys)):
        raise ValueError("观察文件存在重复 Note ID")
    return {
        "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "tables": [
            {
                "name": TABLE_NAME,
                "fields": FIELDS,
                "field_properties": {
                    "观察状态": {
                        "options": [
                            {"name": "相关", "color": 0},
                            {"name": "待复核", "color": 1},
                        ]
                    }
                },
                "hidden_fields": ["来源键"],
                "records": records,
            }
        ],
    }


def import_observations(service, path):
    """幂等导入；不调用小红书、不进入业务 notes 表。"""
    plan = observation_plan(path)
    result = import_plan(service, plan)
    result["views"] = configure_observation_views(service)
    service.repo.set_meta(
        "natural_observations_import",
        str({"source": SOURCE_LABEL, "records": len(plan["tables"][0]["records"])}),
    )
    return {**result, "delivery_excluded": True, "one_time_capture": True}


def configure_observation_views(service):
    """保留两个用途不同的视图，隐藏仅用于幂等导入的来源键。"""
    client = Feishu(service.repo)
    manifest = json.loads(service.repo.meta("project_migration_tables", "{}"))
    table = manifest.get(TABLE_NAME)
    if not table:
        raise ValueError("品牌自然观察表尚未建立")
    path = f"/open-apis/bitable/v1/apps/{client.app}/tables/{table}"
    fields = {field["field_name"]: field for field in client.pages(path + "/fields")}
    views = {
        view["view_name"]: view["view_id"] for view in client.pages(path + "/views")
    }
    if "工作台" in views and "全部观察" not in views:
        client.request(
            "PATCH", path + "/views/" + views["工作台"], json={"view_name": "全部观察"}
        )
        views["全部观察"] = views.pop("工作台")
    if "全部观察" not in views:
        views["全部观察"] = client.request(
            "POST", path + "/views", json={"view_name": "全部观察", "view_type": "grid"}
        )["view"]["view_id"]
    if "待复核" not in views:
        views["待复核"] = client.request(
            "POST", path + "/views", json={"view_name": "待复核", "view_type": "grid"}
        )["view"]["view_id"]
    hidden = [fields["来源键"]["field_id"]]
    client.request(
        "PATCH",
        path + "/views/" + views["全部观察"],
        json={"property": {"hidden_fields": hidden}},
    )
    status = fields["观察状态"]
    options = {item["name"]: item["id"] for item in status["property"]["options"]}
    condition = {
        "field_id": status["field_id"],
        "field_type": status["type"],
        "condition_id": "con"
        + hashlib.sha256(b"natural-observation-review").hexdigest()[:8],
        "operator": "is",
        "value": json.dumps([options["待复核"]], ensure_ascii=False),
    }
    client.request(
        "PATCH",
        path + "/views/" + views["待复核"],
        json={
            "property": {
                "hidden_fields": hidden,
                "filter_info": {"conjunction": "and", "conditions": [condition]},
            }
        },
    )
    return {"all": views["全部观察"], "review": views["待复核"]}
