"""先看交付结论，再看明细和待办；技术底稿默认隐藏。"""

import json
from collections import defaultdict
from datetime import UTC, datetime
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .brands import identity
from .domain import METRICS, month_due
from .presentation import month_views
from .report import LABELS, _build_audit_report, display_time, text
from .urls import validate_url

REPORT_VERSION = "review-v3-contract"
BLUE = "174E63"
TEAL = "D9EDEA"
AMBER = "FFF0CD"


def append(sheet, values):
    sheet.append([text(v) for v in values])


def metric_sum(rows, key, monthly=False):
    values = [r["counts"].get(key) for r in rows if not monthly or r["monthly"]]
    known = [v for v in values if v is not None]
    return sum(known) if known else None


def format_sheet(sheet):
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.font = Font(name="微软雅黑", color="FFFFFF", bold=True, size=11)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for cells in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[cells[0].row].height = 38
        for cell in cells:
            cell.font = Font(name="微软雅黑", size=10, color="243A45")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F7F8")
            if isinstance(cell.value, int):
                cell.number_format = "#,##0"
            if cell.value in ("待确认", "需修改", "拒绝", "待审核"):
                cell.fill = PatternFill("solid", fgColor=AMBER)
    for col, header in enumerate(sheet[1], 1):
        label = header.value or ""
        width = 14
        if any(
            k in label
            for k in ("标题", "事项", "说明", "建议", "原因", "正文", "证据", "关注")
        ):
            width = 42 if "正文" not in label else 65
        elif any(k in label for k in ("时间", "日期", "发布时间")):
            width = 23
        elif "链接" in label or label == "查看原笔记":
            width = 40
        elif any(k in label for k in ("状态", "口径", "结论")):
            width = 25
        sheet.column_dimensions[get_column_letter(col)].width = width
    # 所有业务明细的链接均为最后一列。
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row, sheet.max_column)
        if isinstance(cell.value, str) and cell.value.startswith("https://"):
            try:
                validate_url(cell.value)
            except Exception:
                continue
            cell.hyperlink = cell.value
            cell.font = Font(name="微软雅黑", color="1676B8", underline="single")


def build_review_report(repo, month, output, brand=None):
    month_due(month)
    output = Path(output)
    if output.exists():
        raise ValueError("报表已存在，请选择新的文件名，避免覆盖审核中的文件")
    output.parent.mkdir(parents=True, exist_ok=True)
    rows = month_views(repo, month, brand)
    with TemporaryDirectory(dir=output.parent) as tmp:
        raw = Path(tmp) / "audit.xlsx"
        _build_audit_report(repo, month, raw)
        book = load_workbook(raw)
    audit = book["笔记明细"]
    audit.title = "采集底稿"
    for index in range(audit.max_row, 1, -1):
        nid, submitted = audit.cell(index, 2).value, audit.cell(index, 1).value
        effective, _ = identity(repo, nid, submitted)
        if brand and effective != brand:
            audit.delete_rows(index)
        else:
            audit.cell(index, 1, text(effective))
    for sheet in list(book):
        if sheet != audit:
            book.remove(sheet)
    overview = book.create_sheet("交付总览", 0)
    summary = book.create_sheet("品牌月度汇总", 1)
    detail = book.create_sheet("笔记明细", 2)
    koc = book.create_sheet("KOC汇总", 3)
    performance = book.create_sheet("24小时表现", 4)
    content = book.create_sheet("内容与审核", 5)
    problems = book.create_sheet("异常与待补采", 6)
    notes = book.create_sheet("统计说明", 7)
    overview.append(["审核项目", "结果", "说明"])
    unique = {r["note_id"]: r for r in rows}
    complete = sum(bool(r["monthly"]) for r in unique.values())
    pending = sum(r["review"] != "通过" for r in unique.values())
    contract = json.loads(repo.meta("contract_reviews:" + month, "{}"))
    if repo.meta("workspace_version") == "2":
        pending = sum(
            not contract.get(nid, {}).get("accepted")
            and contract.get(nid, {}).get("status") != "不计入"
            for nid in unique
        )
    unresolved = sum(r["brand"] == "待确认" for r in unique.values())
    unknown = repo.rows(
        "SELECT DISTINCT s.note_id,s.brand,s.url,s.error FROM submissions s LEFT JOIN notes n ON n.note_id=s.note_id WHERE n.published_month IS NULL"
    )
    unknown = [
        r
        for r in unknown
        if not brand or identity(repo, r["note_id"], r["brand"])[0] == brand
    ]
    current_month = month_due(month) > datetime.now(UTC)
    delivery = (
        "本月进行中，尚未月结"
        if current_month
        else "待补齐数据"
        if complete != len(unique) or unresolved or unknown
        else "待审核"
        if pending
        else "审核记录已通过，可准备交付"
    )
    if not unique:
        delivery = "暂无可交付笔记，请先检查异常页"
    if any(r["detail"].get("source") == "mock" for r in rows):
        delivery = "包含演示数据，不可交付"
    for values in [
        ["报告月份", month, "按笔记实际发布时间归属月份"],
        ["交付状态", delivery, "系统不代替你审批或向上发送"],
        ["已登记笔记", len(unique), "去重笔记数；多品牌明细合计可能重复"],
        [
            "涉及品牌",
            len({r["brand"] for r in rows if r["brand"] != "待确认"}),
            f"另有{unresolved}条笔记品牌待确认",
        ],
        [
            "KOC人数",
            len(
                {r["detail"].get("author_id") or r["note_id"] for r in unique.values()}
            ),
            "作者ID去重；作者未知的笔记分别统计",
        ],
        ["月结完成", f"{complete}/{len(unique)}", "交付合计仅使用该月月结快照"],
        [
            "有效24小时数据",
            sum(r["observed_valid"] for r in unique.values()),
            "仅发布后24–26小时的有效数据可比较",
        ],
        ["待审核笔记", pending, "请查看“内容与审核”，审核结论来自人工记录"],
        [
            "月份待确认提交",
            len(unknown),
            "尚未确定月份，列入异常页；不直接算作本月成果",
        ],
        [
            "先看哪里",
            "品牌月度汇总 → 异常与待补采 → 内容与审核",
            "通过原笔记链接可逐条复核",
        ],
    ]:
        append(overview, values)
    summary.append(
        ["品牌", "笔记数", "KOC人数", "月结完成数", "有效24小时数", "待审核数"]
        + ["月结" + LABELS[k] for k in METRICS]
        + ["月结数据说明"]
    )
    groups = defaultdict(list)
    creators = defaultdict(list)
    detail.append(
        [
            "品牌",
            "品牌状态",
            "作者昵称",
            "标题",
            "类型",
            "发布时间",
            "人工审核结论",
            "数据口径",
        ]
        + list(LABELS.values())
        + ["完整互动量", "关注事项", "查看原笔记"]
    )
    content.append(
        [
            "品牌",
            "作者昵称",
            "标题",
            "正文",
            "话题",
            "选题",
            "内容形式",
            "核心卖点",
            "复用建议",
            "品牌识别依据",
            "人工审核结论",
            "审核备注",
            "查看原笔记",
        ]
    )
    problems.append(["品牌", "笔记/提交", "需要处理的事项", "处理建议", "查看原笔记"])
    for r in rows:
        d = r["detail"]
        groups[r["brand"]].append(r)
        creators[
            (
                r["brand"],
                d.get("author_id") or r["note_id"],
                d.get("author_name") or "作者未知",
            )
        ].append(r)
        append(
            detail,
            [
                r["brand"],
                r["brand_status"],
                d.get("author_name"),
                d.get("title"),
                d.get("note_type"),
                display_time(r["published_at"]),
                r["review"],
                r["basis"],
            ]
            + [r["counts"].get(k) for k in METRICS]
            + [r["total"], "；".join(r["issues"]), r["url"]],
        )
        a = r["annotation"]
        append(
            content,
            [
                r["brand"],
                d.get("author_name"),
                d.get("title"),
                d.get("description"),
                d.get("tags"),
                a.get("topic"),
                a.get("format"),
                a.get("selling_points"),
                a.get("reuse_advice"),
                r["brand_evidence"],
                r["review"],
                r["remarks"],
                r["url"],
            ],
        )
        for issue in r["issues"]:
            advice = (
                "历史24小时数据不能补造；本条不参与该阶段排名"
                if "24小时" in issue
                else "核对正文/图片中的品牌，确认后归档"
                if "品牌" in issue
                else "查看原笔记并核对；必要时补采"
            )
            append(
                problems,
                [r["brand"], d.get("title") or r["note_id"], issue, advice, r["url"]],
            )
        if r["review"] in ("需修改", "拒绝"):
            append(
                problems,
                [
                    r["brand"],
                    d.get("title"),
                    r["review"],
                    r["remarks"] or "按人工审核意见处理",
                    r["url"],
                ],
            )
    for name, items in sorted(groups.items()):
        coverage = "；".join(
            f"{LABELS[k]}{sum(r['monthly'] is not None and r['counts'].get(k) is not None for r in items)}/{len(items)}条有值"
            for k in METRICS
        )
        append(
            summary,
            [
                name,
                len(items),
                len({r["detail"].get("author_id") or r["note_id"] for r in items}),
                sum(bool(r["monthly"]) for r in items),
                sum(r["observed_valid"] for r in items),
                sum(r["review"] != "通过" for r in items),
            ]
            + [metric_sum(items, k, monthly=True) for k in METRICS]
            + [coverage],
        )
    koc.append(
        ["品牌", "作者昵称", "笔记数", "月结完成数", "待审核数", "数据口径"]
        + list(LABELS.values())
        + ["数据覆盖说明", "查看原笔记"]
    )
    for (name, _, author), items in sorted(creators.items()):
        append(
            koc,
            [
                name,
                author,
                len(items),
                sum(bool(r["monthly"]) for r in items),
                sum(r["review"] != "通过" for r in items),
                "仅月结数据",
            ]
            + [metric_sum(items, key, monthly=True) for key in METRICS]
            + ["缺失指标不按0统计；每项覆盖条数见品牌汇总", items[0]["url"]],
        )
    performance.append(
        ["排名", "品牌", "作者昵称", "标题", "24小时数据状态"]
        + ["24小时" + LABELS[k] for k in METRICS]
        + ["24小时完整互动量", "选题", "复用建议", "查看原笔记"]
    )
    ranked = sorted(
        rows,
        key=lambda r: (
            r["observed_total"] is None,
            -(r["observed_total"] or 0),
            r["note_id"],
        ),
    )
    rank = 0
    for r in ranked:
        if r["observed_total"] is not None:
            rank += 1
        append(
            performance,
            [
                rank if r["observed_total"] is not None else None,
                r["brand"],
                r["detail"].get("author_name"),
                r["detail"].get("title"),
                r["observed_status"],
            ]
            + [r["observed_counts"].get(k) for k in METRICS]
            + [
                r["observed_total"],
                r["annotation"].get("topic"),
                r["annotation"].get("reuse_advice"),
                r["url"],
            ],
        )
    for r in unknown:
        append(
            problems,
            [
                identity(repo, r["note_id"], r["brand"])[0],
                "月份未确定",
                "尚未取得发布时间或笔记详情",
                "先完成解析/采集，再确认月份归属",
                r["url"],
            ],
        )
    # 仅列入当前报告范围的任务错误，避免其他品牌/月的数据混入。
    scope = {r["note_id"]: r for r in rows}
    for j in repo.rows(
        "SELECT note_id,error,phase FROM jobs WHERE status IN ('FAILED','BLOCKED','RETRY_WAIT')"
    ):
        if j["note_id"] in scope:
            r = scope[j["note_id"]]
            append(
                problems,
                [
                    r["brand"],
                    r["detail"].get("title"),
                    "采集任务需处理",
                    "错误代码见底稿/系统任务；处理后补采，不绕过访问限制",
                    r["url"],
                ],
            )
    notes.append(["项目", "说明"])
    for values in [
        [
            "阅读顺序",
            "先看交付总览，再看品牌汇总、异常待办及内容审核；笔记链接在各明细末列。",
        ],
        ["当前数据", "最近一次公开刷新值；月结完成前仅供跟进，不作为已完成月结成果。"],
        [
            "月结数据",
            "次月1日09:00（北京时间）起实际采集的上月笔记数据；晚采会标注，不推算月末瞬时值。",
        ],
        [
            "24小时表现",
            "发布后24–26小时的有效快照；迟交或晚采不排名，也不拿当前数据冒充历史值。",
        ],
        ["完整互动量", "点赞+收藏+评论+分享；四项必须都有精确值，缺一项则留空。"],
        [
            "曝光与阅读",
            "源站无真实授权数据时留空，空白不是0。人工补录及证据保存在采集底稿，未混入公开数据合计。",
        ],
        [
            "品牌识别",
            "明确正文/话题或图片文字证据才自动归档；多品牌或证据不足的笔记保留待确认。",
        ],
        [
            "审核",
            "AI选题和复用建议不等于人工通过；交付总览的状态来自数据完整性与人工审核记录。",
        ],
        [
            "汇总范围",
            "品牌汇总按品牌关联统计；交付总览按笔记去重。跨品牌重复笔记不能直接累加。",
        ],
        [
            "数据底稿",
            "技术字段与各阶段原始精度保存在隐藏的“采集底稿”工作表，必要时可取消隐藏核查。",
        ],
    ]:
        append(notes, values)
    if repo.meta("workspace_version") == "2":
        accepted = sum(bool(contract.get(nid, {}).get("accepted")) for nid in unique)
        append(
            overview,
            ["合同有效笔记", accepted, "唯一Note ID；仅明确通过且有审核依据的记录计数"],
        )
        for values in overview.iter_rows(min_row=2):
            if values[0].value == "已登记笔记":
                values[0].value = "已收录唯一笔记"
                values[2].value = "按真实Note ID去重；提交次数和短链不计篇数"
            elif values[0].value == "KOC人数":
                values[0].value = "发布账号数"
                values[2].value = "按发布作者ID去重，不代表已确认合作的KOC数"
            elif values[0].value == "待审核笔记":
                values[0].value = "待合同审核"
                values[2].value = "与飞书合同验收一致；未填依据的通过仍待审核"
        koc.title = "发布账号汇总"
        for cell in summary[1]:
            if cell.value == "KOC人数":
                cell.value = "发布账号数"
        for sheet in (detail, content):
            for column in range(sheet.max_column, 0, -1):
                if sheet.cell(1, column).value in ("人工审核结论", "审核备注"):
                    sheet.delete_cols(column)
            last = sheet.max_column
            sheet.insert_cols(last, 2)
            sheet.cell(1, last, "合同验收")
            sheet.cell(1, last + 1, "合同审核依据")
            for i, item in enumerate(rows, 2):
                review = contract.get(item["note_id"], {})
                sheet.cell(i, last, text(review.get("status", "待审核")))
                sheet.cell(i, last + 1, text(review.get("evidence", "")))
    for sheet in book:
        if sheet != audit:
            format_sheet(sheet)
    overview.freeze_panes = "B2"
    overview.column_dimensions["A"].width = 24
    overview.column_dimensions["B"].width = 52
    overview.column_dimensions["C"].width = 75
    notes.column_dimensions["B"].width = 110
    summary.column_dimensions[get_column_letter(summary.max_column)].width = 65
    if performance.max_row > 1:
        performance.conditional_formatting.add(
            f"L2:L{performance.max_row}",
            ColorScaleRule(
                start_type="min",
                start_color="FFFFFF",
                end_type="max",
                end_color="69B7A8",
            ),
        )
    audit.sheet_state = "hidden"
    book.active = 0
    book.properties.title = f"KOC {month} 交付审核"
    book.properties.description = (
        "当前数据与月结数据分开；人工确认后交付；技术底稿隐藏保留。"
    )
    book.save(output)
    output.chmod(0o600)
    result = {
        "path": str(output),
        "notes": len(unique),
        "month": month,
        "sheets": book.sheetnames,
        "delivery_status": delivery,
    }
    book.close()
    return result
