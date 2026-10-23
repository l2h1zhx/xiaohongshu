import json
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .domain import METRICS, SH, month_due, parse_time

LABELS = {
    "likes": "点赞",
    "collects": "收藏",
    "comments": "评论",
    "shares": "分享",
    "exposure": "曝光",
    "reads": "阅读",
}


def text(value):
    # 不让正文、昵称等内容被 Excel 当作公式执行。
    if isinstance(value, str):
        value = "".join(c for c in value if ord(c) >= 32 or c in "\n\t")[:32767]
        if value.lstrip().startswith(("=", "+", "-", "@")):
            value = "'" + value
    return value


def display_time(value):
    return (
        parse_time(value).astimezone(SH).strftime("%Y-%m-%d %H:%M:%S")
        if value
        else None
    )


def complete_total(counts):
    values = [counts.get(k) for k in METRICS[:4]]
    return sum(values) if all(v is not None for v in values) else None


def _build_audit_report(repo, month, output, brand=None):
    month_due(month)
    output = Path(output)
    if output.exists():
        raise ValueError("报表已存在，请选择新的文件名，避免覆盖审核中的文件")
    query = """SELECT n.*,s.brand,min(s.submitted_at) AS submitted_at FROM notes n
               JOIN submissions s ON s.note_id=n.note_id WHERE n.published_month=?"""
    params = [month]
    if brand:
        query += " AND s.brand=?"
        params.append(brand)
    query += " GROUP BY n.note_id,s.brand ORDER BY s.brand,n.published_at,n.note_id"
    rows = repo.rows(query, params)
    book = Workbook()
    book.remove(book.active)
    summary = book.create_sheet("品牌月度汇总")
    detail_sheet = book.create_sheet("笔记明细")
    performance = book.create_sheet("24小时表现")
    problems = book.create_sheet("异常与待补采")
    summary.append(
        [
            "品牌",
            "发布月份",
            "已登记笔记数",
            "月结成功数",
            "月结缺失数",
            "有效24H数",
            "完整互动量可统计笔记数",
            "结案完整互动合计",
            "月结曝光有值数",
            "月结阅读有值数",
            "说明",
        ]
        + [
            label
            for k in METRICS
            for label in ("结案" + LABELS[k] + "合计", LABELS[k] + "有效笔记数")
        ]
    )
    headers = [
        "品牌",
        "笔记ID",
        "作者ID",
        "作者昵称",
        "标题",
        "正文",
        "话题",
        "类型",
        "图片数量",
        "发布时间(北京时间)",
        "登记时间(北京时间)",
        "来源",
        "来源版本",
    ]
    for phase in ("录入", "24H", "结案"):
        headers += [phase + LABELS[k] for k in METRICS] + [
            phase + "采集时间",
            phase + "完整互动量",
            phase + "状态",
            phase + "原始指标与精度",
            phase + "可用字段互动合计",
            phase + "有效互动字段数",
        ]
    headers += [
        "人工补录曝光",
        "人工补录阅读",
        "人工统计时间",
        "人工证据说明",
        "24H相对录入互动增量",
        "内容检查",
        "提醒",
        "AI选题",
        "AI内容形式",
        "AI核心卖点",
        "AI复用建议",
        "AI分析时间",
        "人工审核结论",
        "查看原笔记",
    ]
    detail_sheet.append(headers)
    performance.append(
        [
            "品牌",
            "笔记ID",
            "标题",
            "话题",
            "发布时间",
            "实际24H采集时间",
            "采集时笔记年龄(小时)",
            "24H完整互动量",
            "24H状态",
            "数据来源",
            "查看原笔记",
        ]
    )
    problems.append(["品牌", "笔记ID/提交编号", "异常", "处理建议", "查看原笔记"])
    aggregates = {}
    metric_sums = {}
    for row in rows:
        snapshots = repo.rows(
            "SELECT * FROM snapshots WHERE note_id=? ORDER BY id", (row["note_id"],)
        )
        phases = {}
        manual = None
        for snap in snapshots:
            if (
                snap["phase"] == "MANUAL"
                and json.loads(snap["detail"])["source"] == "manual"
            ):
                manual = snap
            if snap["phase"] in ("INITIAL", "24H") or (
                snap["phase"] == "MONTHLY" and snap["phase_key"] == month
            ):
                phases[snap["phase"]] = snap
        current = json.loads(row["detail"])
        published = row["published_at"]
        output_row = [
            row["brand"],
            row["note_id"],
            current.get("author_id"),
            current.get("author_name"),
            current.get("title"),
            current.get("description"),
            current.get("tags"),
            current.get("note_type"),
            current.get("image_count"),
            display_time(published),
            display_time(row["submitted_at"]),
            current.get("source"),
            current.get("source_version"),
        ]
        totals = {}
        for phase in ("INITIAL", "24H", "MONTHLY"):
            snap = phases.get(phase)
            value = json.loads(snap["detail"]) if snap else {}
            valid = snap is not None and (phase != "24H" or snap["quality"] == "VALID")
            counts = value.get("counts", {}) if valid else {}
            totals[phase] = complete_total(counts)
            state = snap["quality"] if snap else "MISSING"
            if not snap and phase == "MONTHLY" and month_due(month) > datetime.now(UTC):
                state = "NOT_DUE"
            if not snap and phase == "24H":
                missed = repo.rows(
                    "SELECT status FROM jobs WHERE note_id=? AND phase='24H'",
                    (row["note_id"],),
                )
                state = missed[0]["status"] if missed else "UNKNOWN_PUBLISH_TIME"
            output_row += [counts.get(k) for k in METRICS]
            output_row += [
                display_time(snap["captured_at"]) if snap else None,
                totals[phase],
                state,
                json.dumps(
                    {
                        "raw": value.get("raw_counts", {}),
                        "precision": value.get("availability", {}),
                    },
                    ensure_ascii=False,
                ),
                sum(v for k, v in counts.items() if k in METRICS[:4] and v is not None)
                if any(counts.get(k) is not None for k in METRICS[:4])
                else None,
                sum(counts.get(k) is not None for k in METRICS[:4]),
            ]
        manual_detail = json.loads(manual["detail"]) if manual else {}
        annotation = repo.rows(
            "SELECT * FROM annotations WHERE note_id=? ORDER BY id DESC LIMIT 1",
            (row["note_id"],),
        )
        analysis = json.loads(annotation[0]["data"]) if annotation else {}
        reviews = repo.rows(
            "SELECT decision FROM reviews WHERE note_id=? ORDER BY id DESC LIMIT 1",
            (row["note_id"],),
        )
        manual_counts = manual_detail.get("counts", {})
        delta = (
            totals["24H"] - totals["INITIAL"]
            if totals["24H"] is not None and totals["INITIAL"] is not None
            else None
        )
        content = "；".join(
            [
                "标题已读取" if current.get("title") else "标题缺失",
                "正文已读取" if current.get("description") else "正文缺失",
                "图片/视频内容未核验",
                "未提供任务标准时不判定内容达标",
            ]
        )
        output_row += [
            manual_counts.get("exposure"),
            manual_counts.get("reads"),
            display_time(manual["captured_at"]) if manual else None,
            manual_detail.get("evidence"),
            delta,
            content,
            row["warning"],
            analysis.get("topic"),
            analysis.get("format"),
            analysis.get("selling_points"),
            analysis.get("reuse_advice"),
            display_time(annotation[0]["created_at"]) if annotation else None,
            reviews[0]["decision"] if reviews else "待审核",
            row["url"],
        ]
        detail_sheet.append([text(v) for v in output_row])
        observed = phases.get("24H")
        observed_detail = json.loads(observed["detail"]) if observed else {}
        performance.append(
            [
                text(v)
                for v in [
                    row["brand"],
                    row["note_id"],
                    current.get("title"),
                    current.get("tags"),
                    display_time(published),
                    display_time(observed["captured_at"]) if observed else None,
                    round(
                        (
                            parse_time(observed["captured_at"]) - parse_time(published)
                        ).total_seconds()
                        / 3600,
                        2,
                    )
                    if observed
                    else None,
                    totals["24H"],
                    observed["quality"] if observed else "MISSING",
                    observed_detail.get("source"),
                    row["url"],
                ]
            ]
        )
        stats = aggregates.setdefault(row["brand"], [0] * 8)
        stats[0] += 1
        monthly = phases.get("MONTHLY")
        stats[1] += bool(monthly)
        stats[2] += bool(observed and observed["quality"] == "VALID")
        if totals["MONTHLY"] is not None:
            stats[3] += 1
            stats[4] += totals["MONTHLY"]
        month_counts = json.loads(monthly["detail"])["counts"] if monthly else {}
        coverage = metric_sums.setdefault(row["brand"], {k: [0, 0] for k in METRICS})
        for k in METRICS:
            if month_counts.get(k) is not None:
                coverage[k][0] += month_counts[k]
                coverage[k][1] += 1
        stats[5] += month_counts.get("exposure") is not None
        stats[6] += month_counts.get("reads") is not None
        if not monthly and month_due(month) <= datetime.now(UTC):
            problems.append(
                [
                    row["brand"],
                    row["note_id"],
                    "月结数据缺失",
                    "查看任务状态并补采",
                    row["url"],
                ]
            )
        if not observed or observed["quality"] != "VALID":
            problems.append(
                [
                    row["brand"],
                    row["note_id"],
                    "24H数据缺失或晚采",
                    "不可补造历史；晚采原值见历史快照",
                    row["url"],
                ]
            )
        if row["warning"]:
            problems.append(
                [row["brand"], row["note_id"], row["warning"], "人工核对", row["url"]]
            )
    for name, s in aggregates.items():
        summary.append(
            [
                text(name),
                month,
                s[0],
                s[1],
                s[0] - s[1],
                s[2],
                s[3],
                s[4] if s[3] else None,
                s[5],
                s[6],
                "缺失不按0；跨品牌重复笔记不可直接合计；月结为次月实际采集值",
            ]
            + [
                v
                for k in METRICS
                for v in (
                    metric_sums[name][k][0] if metric_sums[name][k][1] else None,
                    metric_sums[name][k][1],
                )
            ]
        )
    # 未解析/发布时间未知的提交不能悄悄从月报消失。
    for row in repo.rows(
        "SELECT s.* FROM submissions s LEFT JOIN notes n ON n.note_id=s.note_id WHERE (s.note_id IS NULL OR n.published_month IS NULL)"
        + (" AND s.brand=?" if brand else ""),
        (brand,) if brand else (),
    ):
        problems.append(
            [
                text(row["brand"]),
                str(row["id"]),
                row["error"] or "发布时间或笔记ID未知，月份归属未确定",
                "补充信息后重新生成报表",
                row["url"],
            ]
        )
    for row in repo.rows(
        "SELECT j.*,n.url FROM jobs j JOIN notes n ON n.note_id=j.note_id WHERE j.status IN ('FAILED','BLOCKED','RETRY_WAIT')"
    ):
        for association in repo.rows(
            "SELECT DISTINCT brand FROM submissions WHERE note_id=?"
            + (" AND brand=?" if brand else ""),
            (row["note_id"], brand) if brand else (row["note_id"],),
        ):
            problems.append(
                [
                    text(association["brand"]),
                    row["note_id"],
                    row["error"],
                    "排查原因后重试；不要绕过访问限制",
                    row["url"],
                ]
            )
    if not rows:
        summary.append(
            [
                brand or "全部品牌",
                month,
                0,
                0,
                0,
                0,
                0,
                None,
                0,
                0,
                "没有已确认属于该月份的笔记；请检查异常页，不能视为投放量为0",
            ]
        )
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="174E63")
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 23
        if sheet != summary:
            sheet.column_dimensions[get_column_letter(sheet.max_column)].width = 45
            for index in range(2, sheet.max_row + 1):
                cell = sheet.cell(index, sheet.max_column)
                if isinstance(cell.value, str) and cell.value.startswith("https://"):
                    from .domain import SourceError
                    from .urls import validate_url

                    try:
                        validate_url(cell.value)
                        cell.hyperlink = cell.value
                        cell.style = "Hyperlink"
                    except SourceError:
                        pass
    book.properties.description = f"生成时间 {datetime.now(UTC).isoformat()}；公开与人工数据分列；可能含 Mock 数据，请核对来源列。"
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    output.chmod(0o600)
    return {
        "path": str(output),
        "notes": len(rows),
        "month": month,
        "sheets": book.sheetnames,
    }


def build_report(repo, month, output, brand=None):
    from .review_report import build_review_report

    return build_review_report(repo, month, output, brand)
