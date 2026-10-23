"""迁移前检查全部工作表；保留行号/公式/链接，不把来源计划当执行结果。"""

import hashlib
import re
from collections import defaultdict
from pathlib import Path
from urllib.parse import urlsplit
from zipfile import ZipFile

from openpyxl import load_workbook

from .urls import HOSTS, extract_urls

DOMAINS = [
    ("风险", "风险与决策"),
    ("任务|派工|里程碑", "项目任务"),
    ("寄样|样品|库存|物流", "寄样与库存"),
    ("预算|费用|佣金|结算|现金", "费用与结算"),
    ("claim|宣称|合规|素材|sku|产品", "产品与合规"),
    ("账号|账户", "账号管理"),
    ("koc|达人|bd|合作", "KOC合作"),
    ("发布|笔记|内容|brief|排期", "笔记管理"),
    ("总控|总览|合同|产能|复盘", "项目总控"),
]


def inspect_workbook(path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("需要不超过20MB的.xlsx原表导出文件")
    with ZipFile(path) as archive:
        if sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("工作簿展开超过100MB")
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    book = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    output, identifiers, links = [], defaultdict(list), defaultdict(list)
    try:
        if len(book.sheetnames) > 100:
            raise ValueError("工作表数量超过100")
        for sheet in book:
            cells, occupied_rows, formulas_without_cache = [], set(), 0
            # 读取真实单元格，不遍历因残留格式膨胀到百万行的空白矩形。
            for cell in list(sheet._cells.values()):
                if cell.value is None and not cell.hyperlink:
                    continue
                value = cell.value
                if hasattr(value, "isoformat"):
                    value = value.isoformat()
                location = f"{sheet.title}!{cell.coordinate}"
                data = {"cell": cell.coordinate, "value": value}
                if cell.data_type == "f":
                    data["formula"] = value
                    data["cached_value"] = cached[sheet.title][cell.coordinate].value
                    if hasattr(data["cached_value"], "isoformat"):
                        data["cached_value"] = data["cached_value"].isoformat()
                    if data["cached_value"] is None:
                        formulas_without_cache += 1
                hyperlink = cell.hyperlink.target if cell.hyperlink else None
                if hyperlink:
                    data["hyperlink"] = hyperlink
                for url in set(
                    extract_urls(str(value or "")) + ([hyperlink] if hyperlink else [])
                ):
                    if url and urlsplit(url).hostname in HOSTS:
                        links[url].append(location)
                for identifier in set(
                    re.findall(
                        r"(?<![A-Za-z0-9])(?:N\d{3,}|KOC-\d{3,}|POST-\d{3,})(?!\d)",
                        str(value or ""),
                    )
                ):
                    identifiers[identifier].append(location)
                cells.append(data)
                occupied_rows.add(cell.row)
            proposed = next(
                (
                    target
                    for pattern, target in DOMAINS
                    if re.search(pattern, sheet.title, re.IGNORECASE)
                ),
                "待核对映射",
            )
            output.append(
                {
                    "sheet": sheet.title,
                    "visibility": sheet.sheet_state,
                    "nonempty_rows_including_headers": len(occupied_rows),
                    "cells": cells,
                    "merged_ranges": [str(r) for r in sheet.merged_cells.ranges],
                    "formulas_without_cache": formulas_without_cache,
                    "proposed_target": proposed,
                    "mapping_requires_review": True,
                }
            )
    finally:
        book.close()
        cached.close()
    return {
        "source_sha256": source_hash,
        "migration_state": "PREFLIGHT_ONLY",
        "sheets": output,
        "identifier_occurrences": dict(identifiers),
        "repeated_identifiers": {k: v for k, v in identifiers.items() if len(v) > 1},
        "xhs_link_occurrences": dict(links),
        "notes_verified": False,
        "contract_acceptance_verified": False,
        "notice": "重复编号可能是跨表引用，不直接删行；短链须解析真实笔记ID后去重；公式不执行，缺少缓存不能当0。",
    }
