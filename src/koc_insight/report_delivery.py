"""月报入口与Excel附件；不向任何人发送消息，不代替人工交付审批。"""

import hashlib
import json
import re
from pathlib import Path
from uuid import uuid4

import httpx

from .domain import SourceError
from .feishu import Feishu
from .presentation import month_views

REPORT_FIELDS = {
    "月份": 1,
    "品牌范围": 1,
    "交付状态": 1,
    "笔记数": 2,
    "KOC数": 2,
    "待审核数": 2,
    "品牌待确认数": 2,
    "月报Excel": 17,
    "说明": 1,
}


def publish_report(
    service, path, month, delivery_status=None, brand=None, retry_upload=False
):
    path = Path(path)
    if (
        path.suffix != ".xlsx"
        or not path.is_file()
        or path.stat().st_size > 20 * 1024 * 1024
    ):
        raise ValueError("请提供不超过20MB的已生成Excel月报")
    if delivery_status is None:
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True)
        try:
            labels = {
                row[0]: row[1]
                for row in book["交付总览"].iter_rows(min_row=2, values_only=True)
            }
            delivery_status = labels["交付状态"]
        finally:
            book.close()
    client = Feishu(service.repo)
    table = service.repo.meta("feishu_report_table_id")
    tables_path = f"/open-apis/bitable/v1/apps/{client.app}/tables"
    if not table:
        matches = [t for t in client.pages(tables_path) if t.get("name") == "月报交付"]
        if matches and not service.repo.meta("feishu_report_create_state"):
            raise SourceError("FEISHU_REPORT_TABLE_NAME_CONFLICT")
        if not matches:
            if service.repo.meta("feishu_report_create_state") == "UNCERTAIN":
                raise SourceError("FEISHU_REPORT_CREATE_UNCERTAIN")
            service.repo.set_meta("feishu_report_create_state", "CREATING")
            if hasattr(service, "checkpoint"):
                service.checkpoint()
            try:
                result = client.request(
                    "POST",
                    tables_path,
                    json={
                        "table": {
                            "name": "月报交付",
                            "fields": [
                                {"field_name": n, "type": t}
                                for n, t in REPORT_FIELDS.items()
                            ],
                        }
                    },
                )
            except SourceError:
                service.repo.set_meta("feishu_report_create_state", "UNCERTAIN")
                raise
            table = result["table_id"]
        elif len(matches) == 1:
            table = matches[0]["table_id"]
        else:
            raise SourceError("FEISHU_REPORT_TABLE_NAME_CONFLICT")
        service.repo.set_meta("feishu_report_table_id", table)
        service.repo.set_meta("feishu_report_create_state", "DONE")
    fingerprint = hashlib.sha256(path.read_bytes()).hexdigest()
    file_key = "feishu_report_file:" + fingerprint
    file_token = service.repo.meta(file_key)
    upload_error = ""
    if not file_token:
        upload_error = service.repo.meta("feishu_report_upload_blocked")
        if retry_upload:
            upload_error = ""
        if not upload_error:
            client.auth()
            try:
                with path.open("rb") as file:
                    response = client.client.post(
                        "/open-apis/drive/v1/medias/upload_all",
                        headers={"Authorization": "Bearer " + client.token},
                        data={
                            "file_name": path.name,
                            "parent_type": "bitable_file",
                            "parent_node": client.app,
                            "size": str(path.stat().st_size),
                        },
                        files={
                            "file": (
                                path.name,
                                file,
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            )
                        },
                    )
                payload = response.json()
            except (httpx.RequestError, ValueError):
                raise SourceError(
                    "FEISHU_REPORT_UPLOAD_NETWORK_ERROR", retryable=True
                ) from None
            if payload.get("code") == 99991672:
                scopes = sorted(
                    set(
                        re.findall(
                            r"(?:drive|base|bitable):[a-z_:]+",
                            str(payload.get("msg", "")),
                        )
                    )
                )
                upload_error = "附件上传权限不足" + (
                    "：" + " / ".join(scopes) if scopes else ""
                )
                service.repo.set_meta("feishu_report_upload_blocked", upload_error)
            else:
                data = client._decode(response).get("data", {})
                file_token = data.get("file_token")
                if not file_token:
                    raise SourceError("FEISHU_REPORT_UPLOAD_EMPTY")
                service.repo.set_meta(file_key, file_token)
                service.repo.set_meta("feishu_report_upload_blocked", "")
    views = month_views(service.repo, month, brand)
    unique = {r["note_id"]: r for r in views}
    contract = json.loads(service.repo.meta("contract_reviews:" + month, "{}"))
    pending = sum(r["review"] != "通过" for r in unique.values())
    if service.repo.meta("workspace_version") == "2":
        pending = sum(
            not contract.get(nid, {}).get("accepted")
            and contract.get(nid, {}).get("status") != "不计入"
            for nid in unique
        )
    fields = {
        "月份": month,
        "品牌范围": brand or "全部品牌",
        "交付状态": delivery_status,
        "笔记数": len(unique),
        "KOC数": len(
            {r["detail"].get("author_id") or r["note_id"] for r in unique.values()}
        ),
        "待审核数": pending,
        "品牌待确认数": sum(r["brand"] == "待确认" for r in unique.values()),
        "说明": upload_error
        or "Excel含交付总览、品牌汇总、笔记明细、KOC汇总、24小时表现、内容审核、异常待办与统计说明；采集底稿隐藏保留。人工确认后再向上交付。",
    }
    if service.repo.meta("workspace_version") == "2":
        fields["发布账号数"] = fields.pop("KOC数")
        fields["已收录笔记"] = fields.pop("笔记数")
        fields["待合同审核"] = fields.pop("待审核数")
        fields["合同有效笔记"] = sum(
            bool(contract.get(nid, {}).get("accepted")) for nid in unique
        )
        fields["说明"] = fields["说明"].replace("KOC汇总", "发布账号汇总")
    if file_token:
        fields["月报Excel"] = [{"file_token": file_token}]
    else:
        fields["月报Excel"] = []
    record_key = "feishu_report_record:" + month + ":" + (brand or "all")
    record_id = service.repo.meta(record_key)
    base = tables_path + "/" + table + "/records"
    if not record_id:
        result = client.request(
            "POST",
            base + "/search",
            json={
                "filter": {
                    "conjunction": "and",
                    "conditions": [
                        {"field_name": "月份", "operator": "is", "value": [month]},
                        {
                            "field_name": "品牌范围",
                            "operator": "is",
                            "value": [brand or "全部品牌"],
                        },
                    ],
                }
            },
        )
        matches = result.get("items", [])
        if len(matches) > 1 or result.get("has_more"):
            raise SourceError("FEISHU_REPORT_DUPLICATE_RECORD")
        record_id = matches[0]["record_id"] if matches else ""
    if record_id:
        client.request("PUT", base + "/" + record_id, json={"fields": fields})
    else:
        token_key = record_key + ":client_token"
        token = service.repo.meta(token_key) or str(uuid4())
        service.repo.set_meta(token_key, token)
        if hasattr(service, "checkpoint"):
            service.checkpoint()
        result = client.request(
            "POST", base, params={"client_token": token}, json={"fields": fields}
        )
        record_id = result["record"]["record_id"]
    service.repo.set_meta(record_key, record_id)
    return {
        "table_id": table,
        "excel_available": bool(file_token),
        "attention": upload_error or None,
    }
