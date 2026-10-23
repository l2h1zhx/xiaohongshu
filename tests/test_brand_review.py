import json

import httpx
import pytest
from openpyxl import load_workbook
from test_core import INITIAL, NID, PUBLISHED, URL, ready

from koc_insight.brands import (
    classify,
    finding,
    identify_pending,
    identity,
    validate_image_url,
)
from koc_insight.domain import Detail, SourceError, stamp
from koc_insight.feishu import FIELD_TYPES, MAIN_FIELDS, Feishu
from koc_insight.report import build_report


def capture(service, description, tags="", images=None):
    service.ingest(URL, now=INITIAL)

    class Source:
        def collect(self, url):
            return Detail(
                note_id=NID,
                description=description,
                tags=tags,
                image_urls=images or [],
                published_at=stamp(PUBLISHED),
                counts={"likes": 5, "collects": 2, "comments": 1, "shares": 0},
            )

    service.run(Source(), now=INITIAL, min_interval=0)


def test_link_only_recognizes_evidence_without_rekeying_history(system):
    capture(system, "ALPHABRAND 随身装，ALPHABRAND 的包装很方便")
    assert identity(system.repo, NID, "__AUTO__") == ("ALPHABRAND", "自动识别")
    assert system.ingest(URL, now=INITIAL)["duplicate"] == 1
    assert system.repo.rows("SELECT brand FROM submissions")[0]["brand"] == "__AUTO__"
    assert len(system.repo.rows("SELECT * FROM snapshots")) == 1


def test_ambiguous_and_missing_brands_are_not_guessed(system):
    capture(system, "ALPHA ALPHA 与 BETA BETA 对比，SPF SPF 防晒")
    assert finding(system.repo, NID)["status"] == "多品牌待确认"
    assert identity(system.repo, NID, "__AUTO__")[0] == "待确认"
    assert identity(system.repo, NID, "用户指定") == ("用户指定", "已指定")
    with pytest.raises(ValueError, match="证据"):
        classify(system, NID, "OTHER", "ALPHA")


def test_chinese_image_ocr_is_cached_and_does_not_mutate_snapshots(system):
    capture(system, "护肤使用体验", images=["https://sns.xhscdn.com/one"])

    class Reader:
        calls = 0

        def read(self, url):
            self.calls += 1
            return "品牌：测试护肤\n补水润养"

    reader = Reader()
    identify_pending(system, reader)
    identify_pending(system, reader)
    assert reader.calls == 1
    assert finding(system.repo, NID)["brand"] == "测试护肤"
    assert finding(system.repo, NID)["image_status"] == "已检查1/1张"
    assert len(system.repo.rows("SELECT * FROM snapshots")) == 1


def test_ocr_failure_keeps_metrics_and_records_actionable_gap(system):
    capture(system, "使用心得", images=["https://sns.xhscdn.com/one"])

    class Reader:
        def read(self, url):
            raise SourceError("IMAGE_UNAVAILABLE")

    identify_pending(system, Reader())
    item = finding(system.repo, NID)
    assert item["status"] == "待确认"
    assert "IMAGE_UNAVAILABLE" in item["image_status"]
    assert (
        json.loads(system.repo.rows("SELECT detail FROM notes")[0]["detail"])["counts"][
            "likes"
        ]
        == 5
    )


def test_workspace_network_failure_gets_only_one_github_retry(system, monkeypatch):
    monkeypatch.delenv("GITHUB_ACTIONS", raising=False)
    capture(system, "体验分享", images=["https://sns.xhscdn.com/one"])

    class Reader:
        calls = 0

        def read(self, url):
            self.calls += 1
            raise httpx.ConnectError("test network unavailable")

    reader = Reader()
    identify_pending(system, reader)
    assert reader.calls == 1
    monkeypatch.setenv("GITHUB_ACTIONS", "true")
    identify_pending(system, reader)
    identify_pending(system, reader)
    assert reader.calls == 2
    assert finding(system.repo, NID)["method"] == "图片OCR（GitHub）"


@pytest.mark.parametrize(
    "url",
    [
        "https://127.0.0.1/a",
        "https://evil.org/a",
        "https://user:pass@sns.xhscdn.com/a",
        "https://sns.xhscdn.com:8080/a",
    ],
)
def test_image_url_guard_rejects_untrusted_destinations(url):
    with pytest.raises(SourceError):
        validate_image_url(url)


def test_image_http_is_upgraded_before_fetch(monkeypatch):
    monkeypatch.setattr("koc_insight.brands.public_dns", lambda host, **kwargs: None)
    assert validate_image_url("http://sns.xhscdn.com/a") == "https://sns.xhscdn.com/a"


def test_ocr_noise_is_not_promoted_to_a_new_brand(system):
    capture(system, "干皮护理分享", images=["https://sns.xhscdn.com/one"])

    class Reader:
        def read(self, url):
            return "AAA\nAAA\nAAA\nAHA\nAHA\nSPF\nSPF"

    identify_pending(system, Reader())
    assert finding(system.repo, NID)["brand"] == ""
    assert finding(system.repo, NID)["status"] == "待确认"


def test_report_separates_current_from_monthly_and_has_clear_review(system, tmp_path):
    capture(system, "品牌：测试品牌\n补水体验")
    output = tmp_path / "review.xlsx"
    build_report(system.repo, "2026-10", output)
    book = load_workbook(output)
    assert book.active.title == "交付总览"
    assert book["采集底稿"].sheet_state == "hidden"
    detail = book["笔记明细"]
    headers = [c.value for c in detail[1]]
    assert detail.cell(2, headers.index("数据口径") + 1).value == "当前数据（非月结）"
    assert detail.cell(2, headers.index("品牌") + 1).value == "测试品牌"
    assert detail.cell(2, detail.max_column).hyperlink.target == URL
    summary = book["品牌月度汇总"]
    assert summary.cell(2, 7).value is None  # 当前点赞不混入月结合计
    assert book["24小时表现"].cell(2, 1).value is None
    assert "采集时间" not in "".join(headers)
    assert book["内容与审核"].cell(2, 11).value == "待审核"
    book.close()


def test_report_brand_filter_applies_to_hidden_audit(system, tmp_path):
    capture(system, "品牌：测试品牌\n使用心得")
    system.ingest(URL, "另一品牌", now=INITIAL)
    output = tmp_path / "one-brand.xlsx"
    build_report(system.repo, "2026-10", output, brand="测试品牌")
    book = load_workbook(output)
    assert book["笔记明细"].max_row == 2
    assert book["采集底稿"].max_row == 2
    assert book["采集底稿"].cell(2, 1).value == "测试品牌"
    assert "另一品牌" not in str(
        [[c.value for c in row] for sheet in book for row in sheet]
    )
    book.close()


def test_review_table_migration_preserves_old_table_and_hides_technical_fields(
    system, configured, respx_mock
):
    ready(system)
    root = "https://open.feishu.cn/open-apis"
    respx_mock.post(root + "/auth/v3/tenant_access_token/internal").respond(
        200, json={"tenant_access_token": "test-token", "expire": 7200}
    )
    tables = root + "/bitable/v1/apps/test-base/tables"
    respx_mock.get(tables).respond(200, json={"data": {"items": []}})
    create = respx_mock.post(tables).respond(
        200, json={"data": {"table_id": "review-table"}}
    )
    base = tables + "/review-table"
    ids = {n: "f" + str(i) for i, n in enumerate(FIELD_TYPES)}
    respx_mock.get(base + "/fields").respond(
        200,
        json={
            "data": {
                "items": [
                    {"field_id": ids[n], "field_name": n, "type": t}
                    for n, t in FIELD_TYPES.items()
                ]
            }
        },
    )
    respx_mock.get(base + "/views").respond(
        200, json={"data": {"items": [{"view_name": "笔记总览", "view_id": "main"}]}}
    )
    respx_mock.post(base + "/views").respond(
        200, json={"data": {"view": {"view_id": "other"}}}
    )
    main = respx_mock.patch(base + "/views/main").respond(200, json={})
    respx_mock.patch(base + "/views/other").respond(200, json={})
    client = Feishu(system.repo)
    old_key = client.business_key(NID, "A")
    assert client.upgrade_review_layout()["view_id"] == "main"
    client.upgrade_review_layout()
    assert create.call_count == 1
    assert system.repo.meta("feishu_archive_table_id") == "test-table"
    assert client.business_key(NID, "A") != old_key
    hidden = json.loads(main.calls[0].request.content)["property"]["hidden_fields"]
    assert ids["最新采集时间"] in hidden
    assert ids["录入点赞"] in hidden
    assert all(ids[name] not in hidden for name in MAIN_FIELDS)
    created_fields = json.loads(create.calls[0].request.content)["table"]["fields"]
    assert created_fields[0]["field_name"] == "品牌"
    assert created_fields[-1]["field_name"] == "原笔记链接"
    assert Feishu(system.repo).table == "review-table"  # 旧环境变量不拉回旧表


def test_feishu_does_not_clear_manual_review_without_new_user_decision(
    system, configured
):
    ready(system)
    fields = Feishu(system.repo).fields(NID, "A")
    assert "审核结论" not in fields
    assert "审核备注" not in fields
