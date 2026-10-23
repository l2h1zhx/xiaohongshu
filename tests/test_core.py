import json
import sqlite3
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook

from koc_insight.collector import MockCollector, XHSCollector, parse_upstream
from koc_insight.domain import SH, Detail, SourceError, exact_count, month_due, stamp
from koc_insight.report import build_report
from koc_insight.repository import Repository
from koc_insight.service import Service
from koc_insight.urls import Resolver, extract_urls, validate_url

NID = "a" * 24
URL = "https://www.xiaohongshu.com/explore/" + NID
PUBLISHED = datetime(2026, 10, 15, 10, tzinfo=SH)
INITIAL = PUBLISHED + timedelta(hours=1)


def ready(s, now=INITIAL):
    s.ingest(URL, "A", now=now)
    collector = MockCollector(stamp(PUBLISHED))
    s.run(collector, now=now, min_interval=0)
    return collector


@pytest.mark.parametrize(
    "value,expected",
    [
        (None, (None, "missing")),
        (-1, (None, "missing")),
        ("", (None, "missing")),
        (0, (0, "exact")),
        ("12,345", (12345, "exact")),
        ("1.2万", (None, "approximate")),
        ("3万+", (None, "approximate")),
        ("99+", (None, "approximate")),
        (True, (None, "missing")),
        ("bad", (None, "invalid")),
    ],
)
def test_exact_count(value, expected):
    assert exact_count(value) == expected


@pytest.mark.parametrize(
    "url",
    [
        "http://www.xiaohongshu.com/explore/" + NID,
        "https://127.0.0.1/a",
        "https://www.xiaohongshu.com.evil.org/explore/" + NID,
        "https://user:secret@www.xiaohongshu.com/explore/" + NID,
        "https://www.xiaohongshu.com:123/explore/" + NID,
        "https://www.xiaohongshu.com/user/profile/a",
        "https://xhslink.com/",
        "https://www.xiaohongshu.com/explore/" + NID + "\n",
    ],
)
def test_reject_url(url):
    with pytest.raises(SourceError):
        validate_url(url)


def test_extract_markdown_and_chinese():
    assert extract_urls(f"[查看]({URL})、{URL}，已发") == [URL, URL]


def test_safe_redirect(respx_mock):
    short = "https://xhslink.com/o/test"
    respx_mock.get(short).respond(302, headers={"location": URL + "?xsec_token=keep"})
    assert Resolver(check_dns=lambda _: None).resolve(short).endswith("xsec_token=keep")


def test_private_redirect_rejected(respx_mock):
    short = "https://xhslink.com/o/test"
    respx_mock.get(short).respond(
        302, headers={"location": "http://169.254.169.254/latest"}
    )
    with pytest.raises(SourceError, match="INVALID_URL"):
        Resolver(check_dns=lambda _: None).resolve(short)


def test_limited_redirects(respx_mock):
    short = "https://xhslink.com/o/test"
    route = respx_mock.get(short).respond(302, headers={"location": short})
    with pytest.raises(SourceError, match="TOO_MANY_REDIRECTS"):
        Resolver(check_dns=lambda _: None).resolve(short)
    assert route.call_count == 6


def test_initial_and_duplicate(system):
    c = ready(system)
    assert system.ingest(URL, "A", now=INITIAL)["duplicate"] == 1
    system.run(c, now=INITIAL, min_interval=0)
    assert c.calls == 1
    assert len(system.repo.rows("SELECT * FROM snapshots")) == 1


def test_cross_brand_shared(system):
    c = ready(system)
    system.ingest(URL, "B", now=INITIAL)
    system.run(c, now=INITIAL, min_interval=0)
    assert c.calls == 1
    assert len(system.repo.rows("SELECT * FROM notes")) == 1
    assert len(system.repo.rows("SELECT * FROM outbox")) == 2


def test_short_long_share_identity(system):
    ready(system)
    system.ingest("https://xhslink.com/o/test", "A", now=INITIAL)

    class Resolved:
        def resolve(self, url):
            return URL

    system.resolve_pending(Resolved(), INITIAL)
    assert len(system.repo.rows("SELECT * FROM notes")) == 1
    assert len(system.repo.rows("SELECT * FROM submissions")) == 2


def test_24h_not_early(system):
    c = ready(system)
    system.run(
        c, now=PUBLISHED + timedelta(hours=24) - timedelta(seconds=1), min_interval=0
    )
    assert c.calls == 1
    system.run(c, now=PUBLISHED + timedelta(hours=24, minutes=30), min_interval=0)
    snap = system.repo.rows("SELECT * FROM snapshots WHERE phase='24H'")[0]
    assert snap["quality"] == "VALID"
    assert snap["captured_at"] == stamp(PUBLISHED + timedelta(hours=24, minutes=30))


def test_26h_boundary(system):
    c = ready(system)
    system.run(c, now=PUBLISHED + timedelta(hours=26), min_interval=0)
    assert (
        system.repo.rows("SELECT * FROM snapshots WHERE phase='24H'")[0]["quality"]
        == "VALID"
    )


def test_late_submission_missed(system):
    ready(system, PUBLISHED + timedelta(hours=25))
    assert (
        system.repo.rows("SELECT * FROM jobs WHERE phase='24H'")[0]["status"]
        == "MISSED"
    )
    assert not system.repo.rows("SELECT * FROM snapshots WHERE phase='24H'")


def test_late_restart(system):
    c = ready(system)
    system.run(c, now=PUBLISHED + timedelta(days=2), min_interval=0)
    assert (
        system.repo.rows("SELECT * FROM snapshots WHERE phase='24H'")[0]["quality"]
        == "LATE"
    )


def test_month_catchup_and_idempotency(system):
    c = ready(system)
    now = datetime(2026, 11, 2, 12, tzinfo=SH)
    system.run(c, now=now, min_interval=0)
    system.run(c, now=now, min_interval=0)
    assert len(system.repo.rows("SELECT * FROM jobs WHERE phase='MONTHLY'")) == 1
    assert len(system.repo.rows("SELECT * FROM snapshots WHERE phase='MONTHLY'")) == 1
    assert (
        system.repo.rows("SELECT * FROM snapshots WHERE phase='MONTHLY'")[0]["quality"]
        == "LATE"
    )


def test_late_month_enrollment(system):
    c = ready(system, datetime(2026, 11, 3, 12, tzinfo=SH))
    assert c.calls == 2
    assert (
        system.repo.rows("SELECT * FROM snapshots WHERE phase='MONTHLY'")[0]["quality"]
        == "LATE_ENROLLMENT"
    )


def test_month_end_phases_independent(system):
    published = datetime(2026, 10, 31, 23, tzinfo=SH)
    system.ingest(URL, "A", now=published)
    c = MockCollector(stamp(published))
    system.run(c, now=published, min_interval=0)
    system.run(c, now=datetime(2026, 11, 1, 9, tzinfo=SH), min_interval=0)
    assert not system.repo.rows("SELECT * FROM snapshots WHERE phase='24H'")
    system.run(c, now=published + timedelta(hours=24, minutes=10), min_interval=0)
    assert (
        len(
            system.repo.rows("SELECT * FROM snapshots WHERE phase IN ('MONTHLY','24H')")
        )
        == 2
    )


def test_manual_does_not_replace_phases(system):
    c = ready(system)
    system.manual(NID, "request-1", INITIAL)
    system.manual(NID, "request-1", INITIAL)
    system.run(c, now=INITIAL, min_interval=0)
    assert len(system.repo.rows("SELECT * FROM snapshots WHERE phase='MANUAL'")) == 1
    assert len(system.repo.rows("SELECT * FROM snapshots WHERE phase='INITIAL'")) == 1


def test_manual_exposure_provenance(system):
    ready(system)
    d = Detail(
        note_id=NID,
        counts={"exposure": 123, "reads": 100},
        source="manual",
        evidence="用户后台截图；统计时间见 captured_at",
    )
    system.import_manual(d, INITIAL, request_id="evidence-1")
    system.import_manual(d, INITIAL, request_id="evidence-1")
    assert len(system.repo.rows("SELECT * FROM snapshots WHERE phase='MANUAL'")) == 1
    base = json.loads(system.repo.rows("SELECT detail FROM notes")[0]["detail"])
    assert base["source"] == "mock"
    assert base["counts"]["exposure"] is None


def test_immutable_snapshots(system):
    ready(system)
    with pytest.raises(sqlite3.IntegrityError):
        system.repo.db.execute("UPDATE snapshots SET quality='VALID'")
    with pytest.raises(sqlite3.IntegrityError):
        system.repo.db.execute("DELETE FROM snapshots")


def test_retry_and_circuit(system):
    system.ingest(URL, "A", now=INITIAL)

    class Broken:
        def collect(self, url):
            raise SourceError("TIMEOUT", retryable=True)

    system.run(Broken(), now=INITIAL, min_interval=0)
    job = system.repo.rows("SELECT * FROM jobs")[0]
    assert job["status"] == "RETRY_WAIT"
    system.run(Broken(), now=INITIAL + timedelta(minutes=5), min_interval=0)
    system.run(Broken(), now=INITIAL + timedelta(minutes=35), min_interval=0)
    assert system.repo.rows("SELECT * FROM jobs")[0]["status"] == "FAILED"


def test_access_restriction_stops_queue(system):
    system.ingest(
        URL + "\nhttps://www.xiaohongshu.com/explore/" + "b" * 24, "A", now=INITIAL
    )

    class Restricted:
        calls = 0

        def collect(self, url):
            self.calls += 1
            raise SourceError("ACCESS_RESTRICTED", block=True)

    c = Restricted()
    system.run(c, now=INITIAL, min_interval=0)
    system.run(c, now=INITIAL, min_interval=0)
    assert c.calls == 1
    assert system.status()["circuit"] == "ACCESS_RESTRICTED"


def test_leases_and_stale_worker(system):
    system.ingest(URL, "A", now=INITIAL)
    old = system._claim(INITIAL)
    new = system._claim(INITIAL + timedelta(minutes=11))
    detail = MockCollector().collect(URL)
    with pytest.raises(SourceError, match="LEASE_LOST"):
        system._save(old, detail, INITIAL + timedelta(minutes=12))
    system._save(new, detail, INITIAL + timedelta(minutes=12))
    assert len(system.repo.rows("SELECT * FROM snapshots")) == 1


def test_restart_and_backup(system, tmp_path):
    c = ready(system)
    backup = tmp_path / "backup.sqlite3"
    system.repo.backup(backup)
    restored = Repository(backup)
    try:
        Service(restored).run(c, now=INITIAL, min_interval=0)
        assert len(restored.rows("SELECT * FROM snapshots")) == 1
        assert restored.db.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    finally:
        restored.db.close()


def test_report_links_missing_and_formula_safety(system, tmp_path):
    c = ready(system)
    system.run(c, now=PUBLISHED + timedelta(hours=27), min_interval=0)
    system.run(c, now=datetime(2026, 11, 1, 9, tzinfo=SH), min_interval=0)
    detail = json.loads(system.repo.rows("SELECT detail FROM notes")[0]["detail"])
    detail["title"] = '=HYPERLINK("https://evil","click")'
    system.repo.db.execute("UPDATE notes SET detail=?", (json.dumps(detail),))
    target = tmp_path / "report.xlsx"
    result = build_report(system.repo, "2026-10", target)
    assert result["notes"] == 1
    book = load_workbook(target)
    assert len(book.sheetnames) == 9
    assert book.sheetnames[0] == "交付总览"
    assert book["采集底稿"].sheet_state == "hidden"
    sheet = book["笔记明细"]
    headers = [c.value for c in sheet[1]]
    assert "采集时间" not in "".join(headers)
    assert sheet.cell(2, headers.index("曝光") + 1).value is None
    assert sheet.cell(2, headers.index("数据口径") + 1).value == "月结数据"
    audit = book["采集底稿"]
    audit_headers = [c.value for c in audit[1]]
    assert audit.cell(2, audit_headers.index("24H点赞") + 1).value is None
    assert sheet.cell(2, headers.index("标题") + 1).data_type != "f"
    assert sheet.cell(2, sheet.max_column).hyperlink.target == URL
    book.close()


def test_report_unknown_time_not_hidden(system, tmp_path):
    system.ingest("https://xhslink.com/o/test", "A", now=INITIAL)
    path = tmp_path / "empty.xlsx"
    build_report(system.repo, "2026-10", path)
    book = load_workbook(path)
    assert book["异常与待补采"].max_row == 2
    book.close()


def test_wrong_note_id_rejected(system):
    system.ingest(URL, "A", now=INITIAL)

    class Wrong:
        def collect(self, url):
            return Detail(note_id="b" * 24)

    system.run(Wrong(), now=INITIAL, min_interval=0)
    assert not system.repo.rows("SELECT * FROM snapshots")
    assert system.repo.rows("SELECT * FROM jobs")[0]["error"] == "NOTE_ID_MISMATCH"


def test_source_contract_and_no_exposure_inference():
    d = parse_upstream(
        {
            "data": {
                "作品ID": NID,
                "时间戳": PUBLISHED.timestamp() * 1000,
                "点赞数量": "1.2万",
                "分享数量": "-1",
                "曝光": 999,
            }
        }
    )
    assert d.counts["likes"] is None
    assert d.raw_counts["likes"] == "1.2万"
    assert d.counts["shares"] is None
    assert d.counts["exposure"] is None
    assert d.published_at == stamp(PUBLISHED)


def test_empty_upstream_is_not_deleted():
    with pytest.raises(SourceError, match="UPSTREAM_EMPTY_UNCLASSIFIED") as e:
        parse_upstream({"data": None})
    assert e.value.block


def test_readonly_api_body(monkeypatch, respx_mock):
    monkeypatch.setenv("XHS_READONLY_REVIEWED", "true")
    monkeypatch.setenv("XHS_API_TOKEN", "test-only-not-a-secret")
    route = respx_mock.post("http://127.0.0.1:5556/xhs/detail").respond(
        200, json={"data": {"作品ID": NID}}
    )
    XHSCollector().collect(URL)
    body = json.loads(route.calls[0].request.content)
    assert body == {"url": URL, "download": False, "check_record": False}


def test_security_gate(monkeypatch):
    monkeypatch.delenv("XHS_READONLY_REVIEWED", raising=False)
    with pytest.raises(SourceError, match="XHS_SECURITY_REVIEW_REQUIRED"):
        XHSCollector().collect(URL)


def test_500_monthly_restart(system):
    urls = "\n".join(
        f"https://www.xiaohongshu.com/explore/{i:024x}" for i in range(1, 501)
    )
    system.ingest(urls, "A", now=INITIAL)
    c = MockCollector(stamp(PUBLISHED))
    system.run(c, limit=500, now=INITIAL, min_interval=0)
    later = datetime(2026, 11, 1, 9, tzinfo=SH)
    # 先跳过已过期 24H，跑完后再分批月结；重启复用同一数据库。
    system.run(c, limit=500, now=later, min_interval=0)
    system.run(c, limit=200, now=later, min_interval=0)
    reopened = Repository(system.repo.path)
    try:
        Service(reopened).run(c, limit=500, now=later, min_interval=0)
        Service(reopened).run(c, limit=500, now=later, min_interval=0)
    finally:
        reopened.db.close()
    assert len(system.repo.rows("SELECT * FROM snapshots WHERE phase='MONTHLY'")) == 500
    assert not system.repo.rows(
        "SELECT note_id,count(*) AS n FROM snapshots WHERE phase='MONTHLY' GROUP BY note_id HAVING n>1"
    )


def test_year_rollover_and_timezone():
    assert month_due("2026-12") == datetime(2027, 1, 1, 1, tzinfo=UTC)
    with pytest.raises(ValueError):
        month_due("2026-13")
    with pytest.raises(ValueError):
        stamp(datetime(2026, 1, 1))
