from datetime import datetime

from openpyxl import Workbook

from koc_insight.domain import SH, Detail
from koc_insight.observations import TABLE_NAME, natural_link_plan, observation_plan


def _sheet(book, name, rows):
    sheet = book.create_sheet(name)
    for row in rows:
        sheet.append(row)


def test_observation_plan_separates_related_and_review_notes(tmp_path):
    path = tmp_path / "natural.xlsx"
    book = Workbook()
    book.remove(book.active)
    headers = [
        "筛选结论",
        "笔记ID",
        "标题",
        "发布时间（北京时间）",
        "笔记链接URL",
        "点赞数",
        "正文/描述",
    ]
    _sheet(
        book,
        "相关笔记",
        [
            headers,
            [
                "正文命中",
                "a" * 24,
                "相关笔记",
                "2026-10-01 12:00:00",
                "https://www.xiaohongshu.com/explore/" + "a" * 24,
                3,
                "正文",
            ],
        ],
    )
    _sheet(
        book,
        "待复核",
        [
            headers,
            [
                "标题命中",
                "b" * 24,
                "待复核笔记",
                "2026-10-02 12:00:00",
                "https://www.xiaohongshu.com/explore/" + "b" * 24,
                None,
                "正文",
            ],
        ],
    )
    book.save(path)

    plan = observation_plan(path)
    table = plan["tables"][0]
    assert table["name"] == TABLE_NAME
    assert [row["观察状态"] for row in table["records"]] == ["相关", "待复核"]
    assert table["records"][0]["点赞"] == 3
    assert "点赞" not in table["records"][1]
    assert table["records"][0]["数据口径"].endswith("不纳入项目交付")


def test_observation_plan_rejects_duplicate_note_ids(tmp_path):
    path = tmp_path / "duplicates.xlsx"
    book = Workbook()
    book.remove(book.active)
    headers = ["筛选结论", "笔记ID", "标题", "发布时间（北京时间）", "笔记链接URL"]
    row = [
        "正文命中",
        "a" * 24,
        "标题",
        "2026-10-01 12:00:00",
        "https://www.xiaohongshu.com/explore/" + "a" * 24,
    ]
    _sheet(book, "相关笔记", [headers, row])
    _sheet(book, "待复核", [headers, row])
    book.save(path)

    try:
        observation_plan(path)
    except ValueError as error:
        assert "重复 Note ID" in str(error)
    else:
        raise AssertionError("duplicate Note ID should be rejected")


def test_natural_link_plan_is_one_time_and_excluded_from_delivery():
    detail = Detail(
        note_id="a" * 24,
        author_id="b" * 24,
        author_name="自然作者",
        title="品牌自然笔记",
        description="正文",
        note_type="图文",
        tags="品牌",
        published_at=datetime(2026, 10, 26, 12, tzinfo=SH).isoformat(),
        counts={"likes": 3, "collects": 2},
    )
    plan = natural_link_plan(
        [{"detail": detail, "link": "https://www.xiaohongshu.com/explore/" + "a" * 24}]
    )
    record = plan["tables"][0]["records"][0]
    assert record["作者ID"] == "b" * 24
    assert record["点赞"] == 3
    assert "自动刷新" in record["数据口径"]
    assert "月报" in record["关联说明"]
