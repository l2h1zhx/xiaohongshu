import json

import pytest

from koc_insight.domain import SourceError
from koc_insight.profiles import collect_profiles, parse_directory
from koc_insight.urls import Resolver, validate_url

UID = "a" * 24
NID = "b" * 24
URL = "https://www.xiaohongshu.com/user/profile/" + UID


def page(cards):
    return (
        "<script>window.__INITIAL_STATE__="
        + json.dumps({"user": {"notes": [cards]}})
        + "</script>"
    )


def test_profile_input_is_separate_from_note_detail(system):
    with pytest.raises(SourceError):
        validate_url(URL)
    system.ingest(URL)
    assert (
        system.repo.rows("SELECT status FROM submissions")[0]["status"]
        == "PROFILE_REGISTERED"
    )
    assert not system.repo.rows("SELECT * FROM notes")
    assert Resolver().resolve(URL) == URL
    assert len(json.loads(system.repo.meta("profile_requests"))) == 1


def test_visible_directory_has_real_ids_and_no_suggested_authors(system):
    cards = [
        {"id": NID, "xsecToken": "token", "noteCard": {"user": {"userId": UID}}},
        {"id": "c" * 24, "noteCard": {"user": {"userId": "d" * 24}}},
    ]
    links = parse_directory(page(cards), UID)
    assert len(links) == 1 and NID in links[0]
    system.ingest(URL)
    result = collect_profiles(system, reader=lambda _: links)
    assert result["discovered_links"] == 1
    entry = next(iter(json.loads(system.repo.meta("profile_requests")).values()))
    assert entry["status"] == "PARTIAL"
    assert not system.repo.meta("account_coverage:" + UID)
    assert (
        collect_profiles(system, reader=lambda _: pytest.fail("unexpected retry"))[
            "processed"
        ]
        == 0
    )


@pytest.mark.parametrize(
    "html,code",
    [
        ("IP at risk 300012", "PROFILE_ACCESS_RESTRICTED"),
        (page([{"id": "", "noteCard": {}}]), "PROFILE_NO_READABLE_NOTE_IDS"),
        ("<html>请登录</html>", "PROFILE_DIRECTORY_UNAVAILABLE"),
    ],
)
def test_restricted_or_empty_directory_is_not_zero_notes(html, code):
    with pytest.raises(SourceError) as error:
        parse_directory(html, UID)
    assert error.value.code == code


def test_restriction_stops_other_account_attempts(system):
    system.ingest(URL + "\nhttps://www.xiaohongshu.com/user/profile/" + "c" * 24)

    def blocked(_):
        raise SourceError("PROFILE_ACCESS_RESTRICTED", block=True)

    assert collect_profiles(system, reader=blocked, limit=2)["processed"] == 1
    assert (
        collect_profiles(system, reader=lambda _: pytest.fail("unexpected retry"))[
            "processed"
        ]
        == 0
    )


def test_registered_homepage_is_not_an_unresolved_note_in_export(system, tmp_path):
    from openpyxl import load_workbook

    from koc_insight.report import build_report

    system.ingest(URL)
    system.repo.set_meta("workspace_version", "2")
    path = tmp_path / "empty.xlsx"
    build_report(system.repo, "2026-10", path)
    book = load_workbook(path)
    assert book["异常与待补采"].max_row == 1
    book.close()


def test_profile_login_redirect_is_reported_without_following_login(monkeypatch):
    import httpx

    from koc_insight.profiles import read_directory

    monkeypatch.setattr("koc_insight.profiles.public_dns", lambda _: None)
    seen = []

    def respond(request):
        seen.append(str(request.url))
        return httpx.Response(302, headers={"location": "/login?redirectPath=private"})

    with httpx.Client(transport=httpx.MockTransport(respond)) as client:
        with pytest.raises(SourceError) as error:
            read_directory(URL, client)
    assert error.value.code == "PROFILE_LOGIN_REQUIRED" and error.value.block
    assert seen == [URL]
