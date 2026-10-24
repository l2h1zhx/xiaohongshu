import json
from copy import deepcopy

import pytest
from openpyxl import Workbook

from koc_insight.project_hub import counts, sync_hub
from koc_insight.project_import import import_plan, validate_plan
from koc_insight.project_sources import inspect_workbook


def test_workbook_keeps_hidden_sources_formulas_and_references(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "任务"
    sheet.append(["ID", "链接", "合计"])
    sheet.append(
        [
            "N001 POST-001",
            '=HYPERLINK("https://xhslink.com/o/test","查看")',
            "=SUM(D2:E2)",
        ]
    )
    hidden = book.create_sheet("历史")
    hidden.sheet_state = "hidden"
    hidden.append(["N001", None])
    hidden["B2"].hyperlink = "https://xhslink.cn/m/account"
    sheet["Z10000"].number_format = "0.00"
    path = tmp_path / "source.xlsx"
    book.save(path)
    result = inspect_workbook(path)
    assert len(result["sheets"]) == 2
    assert result["sheets"][0]["nonempty_rows_including_headers"] == 2
    assert result["sheets"][0]["formulas_without_cache"] == 2
    assert result["sheets"][1]["visibility"] == "hidden"
    assert result["repeated_identifiers"]["N001"] == ["任务!A2", "历史!A1"]
    assert len(result["xhs_link_occurrences"]) == 2
    assert result["notes_verified"] is result["contract_acceptance_verified"] is False


def test_counts_require_source_scope_and_latest_review(system):
    account = "b" * 24
    for n, month, author, source in [
        ("a", "2026-10", account, "xhs-downloader"),
        ("b", "2026-10", account, "mock"),
        ("c", "2026-09", account, "xhs-downloader"),
        ("d", "2026-10", "e" * 24, "xhs-downloader"),
    ]:
        with system.repo.transaction() as db:
            db.execute(
                "INSERT INTO notes(note_id,url,first_submitted,published_month,detail) VALUES(?,?,?,?,?)",
                (
                    n * 24,
                    "https://example.invalid",
                    "now",
                    month,
                    json.dumps({"author_id": author, "source": source}),
                ),
            )
    with system.repo.transaction() as db:
        db.execute(
            "INSERT INTO reviews(note_id,decision,remarks,created_at) VALUES(?,?,?,?)",
            ("a" * 24, "通过", "", "now"),
        )
        db.execute(
            "INSERT INTO reviews(note_id,decision,remarks,created_at) VALUES(?,?,?,?)",
            ("a" * 24, "需修改", "", "later"),
        )
    result = counts(
        system.repo, {"month": "2026-10", "authorized_account_ids": [account]}
    )
    assert result == {
        "已采集笔记": 1,
        "人工审核通过": 0,
        "待审核笔记": 0,
        "其他审核状态笔记": 1,
    }
    assert counts(system.repo, {"month": "2026-10"})["已采集笔记"] is None


def sample_plan():
    return {
        "source_sha256": "a" * 64,
        "tables": [
            {
                "name": "迁移任务",
                "fields": {"编号": 1, "来源键": 1},
                "records": [{"编号": "N001", "来源键": "source:1"}],
            }
        ],
    }


def test_plan_rejects_duplicate_key_and_bad_field_before_writes():
    plan = sample_plan()
    plan["tables"][0]["records"] *= 2
    with pytest.raises(ValueError):
        validate_plan(plan)
    plan = sample_plan()
    plan["tables"][0]["records"][0]["undeclared"] = "bad"
    with pytest.raises(ValueError):
        validate_plan(plan)


def test_import_is_idempotent_and_keeps_manual_changes(system, monkeypatch):
    records = []
    calls = []

    class FakeFeishu:
        app = "base"

        def __init__(self, repo):
            pass

        def pages(self, path):
            if path.endswith("/fields"):
                return [
                    {"field_name": n, "field_id": n, "type": 1}
                    for n in ("编号", "来源键")
                ]
            if path.endswith("/records"):
                return records
            return []

        def request(self, method, path, **kwargs):
            calls.append((method, path))
            if path.endswith("/tables"):
                return {"table_id": "table"}
            if path.endswith("/batch_create"):
                records.extend(
                    {"record_id": str(i), "fields": deepcopy(r["fields"])}
                    for i, r in enumerate(kwargs["json"]["records"])
                )
                return {"records": records}
            raise AssertionError(path)

    monkeypatch.setattr("koc_insight.project_import.Feishu", FakeFeishu)
    first = import_plan(system, sample_plan())
    records[0]["fields"]["编号"] = "手工修改"
    second = import_plan(system, sample_plan())
    assert first["tables"][0]["created_rows"] == 1
    assert second["tables"][0]["created_rows"] == 0
    assert records[0]["fields"]["编号"] == "手工修改"
    assert len(calls) == 2


def test_hub_refuses_switching_existing_project(system):
    system.repo.set_meta("active_project", "original")
    with pytest.raises(ValueError):
        sync_hub(
            system, {"project_id": "replacement", "name": "new", "month": "2026-10"}
        )
    assert system.repo.meta("active_project") == "original"


def test_relation_links_unique_exact_ids_without_guessing():
    from koc_insight.project_import import relation_values

    target = [
        {"record_id": "one", "fields": {"ID": "N001"}},
        {"record_id": "two", "fields": {"ID": "N002"}},
        {"record_id": "three", "fields": {"ID": "N002"}},
    ]
    source = [
        {"record_id": "a", "fields": {"关联编号": "N001、N002、N0010"}},
        {"record_id": "one", "fields": {"关联编号": "N001"}},
    ]
    assert relation_values(source, target, "ID", r"N\d{3}") == {"a": ["one"]}


def test_feishu_pagination_accepts_empty_null_items(system, monkeypatch):
    from koc_insight.feishu import Feishu

    client = Feishu(system.repo)
    monkeypatch.setattr(
        client, "request", lambda *a, **k: {"items": None, "has_more": False}
    )
    assert list(client.pages("/empty")) == []


def test_legacy_auto_brand_alias_does_not_duplicate_notes(system):
    from koc_insight.brands import AUTO_BRAND
    from koc_insight.feishu import Feishu
    from koc_insight.presentation import month_views

    url = "https://www.xiaohongshu.com/explore/" + "a" * 24
    system.ingest(url, "授权验证（品牌待确认）")
    assert system.ingest(url)["duplicate"] == 1
    with system.repo.transaction() as db:
        db.execute(
            "UPDATE notes SET published_month='2026-10' WHERE note_id=?", ("a" * 24,)
        )
        db.execute(
            "INSERT INTO submissions(brand,url,submitted_at,note_id,status) VALUES(?,?,?,?,?)",
            ("授权验证（品牌待确认）", url + "?old=1", "now", "a" * 24, "REGISTERED"),
        )
    assert len(month_views(system.repo, "2026-10")) == 1
    c = Feishu(system.repo)
    assert c.business_key("a" * 24, AUTO_BRAND) == c.business_key(
        "a" * 24, "授权验证（品牌待确认）"
    )


def test_contract_pass_needs_explicit_evidence_and_reconciles_unique_source(system):
    from koc_insight.workspace import workspace_metrics

    nid = "f" * 24
    with system.repo.transaction() as db:
        db.execute(
            "INSERT INTO notes(note_id,url,first_submitted,published_month,detail) VALUES(?,?,?,?,?)",
            (
                nid,
                "https://example.invalid",
                "now",
                "2026-10",
                json.dumps({"source": "xhs-downloader", "author_id": "a" * 24}),
            ),
        )
    rows = [{"fields": {"Note ID": nid, "合同验收": "通过", "审核依据": ""}}]
    m = workspace_metrics(system.repo, rows, "2026-10")
    assert (
        m["collected"],
        m["accepted"],
        m["pending"],
        m["missing_approval_evidence"],
    ) == (1, 0, 1, 1)
    rows[0]["fields"]["审核依据"] = "负责人核对SKU、原创和披露后确认"
    assert workspace_metrics(system.repo, rows, "2026-10")["accepted"] == 1
    with pytest.raises(Exception, match="WORKSPACE_DUPLICATE_NOTE_ID"):
        workspace_metrics(system.repo, rows * 2, "2026-10")
    with pytest.raises(Exception, match="WORKSPACE_MISSING_COLLECTED_NOTES"):
        workspace_metrics(system.repo, [], "2026-10")


def test_workspace_uses_one_key_across_brand_associations(system):
    from koc_insight.feishu import Feishu

    system.repo.set_meta("workspace_version", "2")
    c = Feishu(system.repo)
    assert c.business_key("a" * 24, "A") == c.business_key("a" * 24, "B")


def test_native_plan_checks_enum_dates_and_attachments():
    plan = sample_plan()
    t = plan["tables"][0]
    t["fields"].update({"状态": 3, "截止": 5, "附件": 17})
    t["records"][0].update(
        {
            "状态": "待审核",
            "截止": 1792700000000,
            "附件": [{"file_token": "test-only-token"}],
        }
    )
    validate_plan(plan)
    t["records"][0]["附件"] = [{"url": "bad"}]
    with pytest.raises(ValueError):
        validate_plan(plan)


def test_workspace_report_keeps_one_contract_count_and_link_last(system, tmp_path):
    from datetime import datetime

    from openpyxl import load_workbook

    from koc_insight.collector import MockCollector
    from koc_insight.domain import SH
    from koc_insight.report import build_report

    nid = "c" * 24
    system.ingest(
        "https://www.xiaohongshu.com/explore/" + nid,
        now=datetime(2026, 10, 15, 11, tzinfo=SH),
    )
    system.run(
        MockCollector(), now=datetime(2026, 10, 15, 11, tzinfo=SH), min_interval=0
    )
    with system.repo.transaction() as db:
        db.execute(
            "INSERT INTO submissions(brand,url,submitted_at,note_id,status) VALUES(?,?,?,?,?)",
            (
                "授权验证（品牌待确认）",
                "https://www.xiaohongshu.com/explore/" + nid + "?old=1",
                "2026-10-15T11:00:00+08:00",
                nid,
                "REGISTERED",
            ),
        )
    system.repo.set_meta("workspace_version", "2")
    system.repo.set_meta(
        "contract_reviews:2026-10",
        json.dumps(
            {
                nid: {
                    "status": "待补证",
                    "accepted": False,
                    "evidence": "测试：待补原文证据",
                }
            }
        ),
    )
    path = tmp_path / "review.xlsx"
    build_report(system.repo, "2026-10", path)
    book = load_workbook(path)
    labels = [r[0].value for r in book["交付总览"].iter_rows(min_row=2)]
    assert book["采集底稿"].max_row == 2
    assert labels.count("待合同审核") == 1
    assert labels.count("合同有效笔记") == 1
    assert "已收录笔记" in labels
    for name in ("笔记明细", "内容与审核"):
        sheet = book[name]
        headers = [c.value for c in sheet[1]]
        assert headers[-1] == "查看原笔记"
        assert "合同验收" in headers and "人工审核结论" not in headers
        assert sheet.cell(2, headers.index("合同验收") + 1).value == "待补证"
    book.close()


def test_workspace_sync_omits_duplicate_stage_columns_and_preserves_contract_fields(
    system, monkeypatch, configured
):
    from datetime import datetime

    from koc_insight.collector import MockCollector
    from koc_insight.domain import SH
    from koc_insight.feishu import Feishu
    from koc_insight.workspace import NOTE_TYPES

    nid = "d" * 24
    system.ingest(
        "https://www.xiaohongshu.com/explore/" + nid,
        now=datetime(2026, 10, 15, 11, tzinfo=SH),
    )
    system.run(
        MockCollector(), now=datetime(2026, 10, 15, 11, tzinfo=SH), min_interval=0
    )
    system.repo.set_meta("workspace_version", "2")
    c = Feishu(system.repo)
    fields = c.fields(nid, "__AUTO__")
    assert (
        "最新点赞" not in fields
        and "录入采集时间" not in fields
        and "结案点赞" not in fields
    )
    assert "合同验收" not in fields and "审核依据" not in fields
    assert fields["归属月份"] == "2026-10"
    monkeypatch.setattr(
        c,
        "pages",
        lambda path: iter(
            [{"field_name": k, "type": v} for k, v in NOTE_TYPES.items()]
        ),
    )
    assert c.diagnostic()["required_fields"] == len(NOTE_TYPES) < 40


def test_new_note_registers_publisher_once_and_links_without_changing_review(
    system, monkeypatch
):
    from koc_insight.workspace import sync_workspace

    nid, uid = "e" * 24, "b" * 24
    with system.repo.transaction() as db:
        db.execute(
            "INSERT INTO notes(note_id,url,first_submitted,published_month,detail) VALUES(?,?,?,?,?)",
            (
                nid,
                "https://example.invalid",
                "now",
                "2026-10",
                json.dumps(
                    {
                        "source": "xhs-downloader",
                        "author_id": uid,
                        "author_name": "测试作者",
                    }
                ),
            ),
        )
    context = {
        "project_id": "test",
        "month": "2026-10",
        "authorized_account_ids": [uid],
    }
    for key, value in {
        "workspace_version": "2",
        "active_project": "test",
        "project_context:test": json.dumps(context),
        "project_migration_tables": json.dumps(
            {"01 项目总控": "hub", "03 账号与达人": "people", "05 笔记与验收": "notes"}
        ),
        "project_hub_record_id": "summary",
    }.items():
        system.repo.set_meta(key, value)
    note = {
        "record_id": "note",
        "fields": {
            "Note ID": nid,
            "合同验收": "待补证",
            "审核依据": "用户要求补原文证据",
            "发布账号": [],
        },
    }
    people = []
    created = []

    class Client:
        app = "base"

        def __init__(self, repo):
            self.repo = repo

        def pages(self, path):
            return people if "/people/" in path else [note]

        def request(self, method, path, **kwargs):
            if method == "GET":
                return {
                    "record": {
                        "fields": {
                            "当前月份": "2026-10",
                            "合同底线": 60,
                            "内部目标": 100,
                        }
                    }
                }
            if path.endswith("/people/records"):
                actor = {
                    "record_id": "author",
                    "fields": deepcopy(kwargs["json"]["fields"]),
                }
                people.append(actor)
                created.append(actor)
                return {"record": deepcopy(actor)}
            if path.endswith("/notes/records/batch_update"):
                note["fields"].update(kwargs["json"]["records"][0]["fields"])
            return {}

    monkeypatch.setattr("koc_insight.workspace.Feishu", Client)
    sync_workspace(system)
    sync_workspace(system)
    assert len(created) == 1 and people[0]["fields"]["本月收录笔记"] == 1
    assert note["fields"]["发布账号"] == ["author"]
    assert (
        note["fields"]["合同验收"] == "待补证"
        and note["fields"]["审核依据"] == "用户要求补原文证据"
    )
    assert people[0]["fields"]["角色"] == "已采集账号（归属待确认）"
